"""
FlowGrid — Dynamic Traffic Congestion Analytics Using Fenwick Tree
PowerPoint Presentation Generator (Updated for Member 4 Integration)
Course: Data Structures & Algorithms (2nd Year CSE Case Study)

Generates the updated, comprehensive 16-slide academic presentation deck
covering the complete 4-member implementation, mathematics, C++ core,
Python HTTP server bridge, web dashboard, and 32-case test suite.
"""

import sys
import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE


def create_deck():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    # -------------------------------------------------------------------------
    # Consistent Dark Academic Color Palette
    # -------------------------------------------------------------------------
    BG_DARK = RGBColor(15, 23, 42)        # Slate 900
    CARD_BG = RGBColor(30, 41, 59)        # Slate 800
    CARD_BORDER = RGBColor(51, 65, 85)    # Slate 700
    TEXT_LIGHT = RGBColor(248, 250, 252)  # Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)  # Slate 400
    ACCENT_BLUE = RGBColor(59, 130, 246)  # Blue 500
    ACCENT_CYAN = RGBColor(6, 182, 212)   # Cyan 500
    ACCENT_GREEN = RGBColor(16, 185, 129) # Emerald 500
    ACCENT_CORAL = RGBColor(239, 68, 68)  # Red 500
    ACCENT_AMBER = RGBColor(245, 158, 11) # Amber 500

    blank_layout = prs.slide_layouts[6]

    def set_slide_background(slide):
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = BG_DARK

    def add_header(slide, title_text, category_text="DSA CASE STUDY EVALUATION"):
        # Category tag
        cat_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(0.35))
        tf_cat = cat_box.text_frame
        tf_cat.word_wrap = True
        p_cat = tf_cat.paragraphs[0]
        p_cat.text = category_text.upper()
        p_cat.font.size = Pt(10)
        p_cat.font.bold = True
        p_cat.font.color.rgb = ACCENT_CYAN

        # Title
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.7), Inches(11.7), Inches(0.7))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        p_title = tf_title.paragraphs[0]
        p_title.text = title_text
        p_title.font.size = Pt(22)
        p_title.font.bold = True
        p_title.font.color.rgb = TEXT_LIGHT

    def add_card(slide, left, top, width, height, title="", border_color=CARD_BORDER):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = CARD_BG
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)

        if title:
            tb = slide.shapes.add_textbox(left + Inches(0.2), top + Inches(0.15), width - Inches(0.4), Inches(0.45))
            tf = tb.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.text = title
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = ACCENT_BLUE

        return shape

    # =========================================================================
    # SLIDE 1: TITLE SLIDE & METADATA
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_slide_background(s1)

    tbox1 = s1.shapes.add_textbox(Inches(1.0), Inches(1.2), Inches(11.3), Inches(3.0))
    tf1 = tbox1.text_frame
    tf1.word_wrap = True

    p0 = tf1.paragraphs[0]
    p0.text = "DATA STRUCTURES & ALGORITHMS (DSA) CASE STUDY"
    p0.font.size = Pt(13)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_CYAN

    p1 = tf1.add_paragraph()
    p1.text = "FlowGrid — Dynamic Traffic Congestion Analytics\nUsing Fenwick Tree (Binary Indexed Tree)"
    p1.font.size = Pt(30)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_LIGHT

    p2 = tf1.add_paragraph()
    p2.text = "Sub-Corridor Range Aggregation & Point Inflow Tracking via 1-Based BIT"
    p2.font.size = Pt(15)
    p2.font.color.rgb = TEXT_MUTED

    # Project metadata card
    add_card(s1, Inches(1.0), Inches(4.5), Inches(11.3), Inches(2.3), "PROJECT & EVALUATION METADATA")
    tbox_team = s1.shapes.add_textbox(Inches(1.2), Inches(5.0), Inches(10.9), Inches(1.6))
    tf_team = tbox_team.text_frame
    tf_team.word_wrap = True

    meta_lines = [
        ("Student Team:", "2nd Year CSE Case Study Team (4-Member Division)"),
        ("Course & Degree:", "Data Structures & Algorithms | 2nd Year B.Tech Computer Science & Engineering"),
        ("Core Data Structure:", "1-Based Fenwick Tree (Binary Indexed Tree / BIT) in C++17"),
        ("Architecture Stack:", "C++17 Engine (traffic_analytics.exe) + Python 3 Bridge (server.py) + Web Dashboard"),
        ("Evaluation Mode:", "Automated Non-Menu Benchmark Verification & Interactive Web Integration")
    ]
    for lbl, val in meta_lines:
        p = tf_team.add_paragraph() if tf_team.paragraphs[0].text else tf_team.paragraphs[0]
        p.text = f"{lbl} {val}"
        p.font.size = Pt(11)
        p.font.color.rgb = ACCENT_GREEN if "Core Data Structure" in lbl else (TEXT_LIGHT if "Architecture" in lbl else TEXT_MUTED)
        p.space_after = Pt(3)

    # =========================================================================
    # SLIDE 2: 4-MEMBER TEAM WORK DIVISION
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_slide_background(s2)
    add_header(s2, "Four-Member Collaborative Architecture & Division of Work", "TEAM RESPONSIBILITIES")

    members = [
        ("Member 1: Data Management & Tree Initialization", [
            "Formulated 1-based indexing representation for 8 corridor segments.",
            "Implemented initializeData() handling raw traffic counts vector.",
            "Engineered O(n) linear-time Fenwick Tree build (buildFenwick()) via forward-propagation technique, avoiding O(n log n) overhead."
        ], Inches(0.8), Inches(1.6)),

        ("Member 2: Core Bitwise & Tree Operations", [
            "Implemented bitwise lowbit(x) = x & (-x) isolating lowest set bit.",
            "Implemented O(log n) point updates: update(idx, delta) climbing tree.",
            "Implemented setCount(idx, newCount) absolute setting with delta calculation.",
            "Implemented O(log n) prefix sum accumulation: prefixSum(idx)."
        ], Inches(6.8), Inches(1.6)),

        ("Member 3: Range Queries & Congestion Detection", [
            "Implemented corridor range sum: rangeSum(L, R) = prefixSum(R) - prefixSum(L - 1).",
            "Implemented threshold classification: detectCongestion(sum, threshold).",
            "Authored non-menu-driven automated C++ demonstration (Steps 1 to 8)."
        ], Inches(0.8), Inches(4.3)),

        ("Member 4: Dashboard Integration, Testing & Parity", [
            "Designed lightweight Python HTTP server (server.py) bridging browser to C++ binary.",
            "Enhanced web dashboard UI (card selection, active range tags, ancestor path tracing).",
            "Standardized threshold logic (< 55 -> NORMAL, >= 55 -> POTENTIALLY CONGESTED).",
            "Engineered comprehensive 32-case automated test suite (test_flowgrid.py)."
        ], Inches(6.8), Inches(4.3))
    ]

    for title, points, left, top in members:
        add_card(s2, left, top, Inches(5.7), Inches(2.5), title)
        tb = s2.shapes.add_textbox(left + Inches(0.2), top + Inches(0.55), Inches(5.3), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        for i, pt in enumerate(points):
            p = tf.add_paragraph() if i > 0 else tf.paragraphs[0]
            p.text = "• " + pt
            p.font.size = Pt(10)
            p.font.color.rgb = TEXT_LIGHT
            p.space_after = Pt(2)

    # =========================================================================
    # SLIDE 3: PROBLEM STATEMENT
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_slide_background(s3)
    add_header(s3, "Problem Statement: Dynamic Urban Traffic Monitoring")

    add_card(s3, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "The Urban Congestion Challenge")
    tb3_1 = s3.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.4))
    tf3_1 = tb3_1.text_frame
    tf3_1.word_wrap = True
    p3_pts = [
        "Metropolitan road corridors experience high volatility in vehicular density across adjoining segments due to rush hours, traffic signals, and incidents.",
        "Road sensor counters report instantaneous volume changes that must be ingested without lagging.",
        "Traffic management authorities need real-time cumulative volume queries over flexible sub-corridors [L..R] to deploy emergency signals and dynamic rerouting.",
        "Threshold rule violations require immediate algorithmic detection to alert operators before physical gridlock occurs."
    ]
    for pt in p3_pts:
        p = tf3_1.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(10)

    add_card(s3, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Algorithmic Inefficiency of Classical Arrays")
    tb3_2 = s3.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.4))
    tf3_2 = tb3_2.text_frame
    tf3_2.word_wrap = True
    p3_pts2 = [
        "Naive Array Approach: Point updates take O(1), but range queries require O(n) linear scans. Severely degraded performance under thousands of queries.",
        "Static Prefix Sum Array: Range queries take O(1), but a single point update cascades across all subsequent indices, taking O(n) time. Unusable for live traffic streams.",
        "The Algorithmic Dilemma: Real-time traffic monitoring requires a balanced data structure supporting BOTH point updates and range queries sub-linearly in O(log n).",
        "Rigorous Constraints: Strict 1-based indexing, minimal auxiliary memory, zero pointer chasing, and robust boundary safety validation."
    ]
    for pt in p3_pts2:
        p = tf3_2.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(12)
        p.font.color.rgb = TEXT_MUTED
        p.space_after = Pt(10)

    # =========================================================================
    # SLIDE 4: OBJECTIVES & SCOPE
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_slide_background(s4)
    add_header(s4, "Project Objectives & Scope of Engineering")

    objs = [
        ("1. Core C++17 Fenwick Tree Engine", "Implement a mathematically verified, standard C++17 Binary Indexed Tree with strict 1-based indexing, boundary safety checks, and an automated non-menu benchmark suite.", Inches(0.8), Inches(1.6)),
        ("2. Dual-Array Architecture", "Maintain clean decoupling between raw traffic sensor counts (trafficCounts[1..n]) and pre-calculated partial sum nodes (bit[1..n]) using strictly O(n) space.", Inches(6.8), Inches(1.6)),
        ("3. Sub-Linear Operations (O(log n))", "Guarantee O(log n) time bounds for point updates (update / setCount) and corridor range queries (rangeSum(L, R)) using lowbit binary decomposition.", Inches(0.8), Inches(4.3)),
        ("4. Full-Stack Dashboard Integration", "Integrate the compiled C++ binary with an interactive web dashboard via a lightweight standard-library Python HTTP server for live demonstration.", Inches(6.8), Inches(4.3))
    ]
    for title, desc, left, top in objs:
        add_card(s4, left, top, Inches(5.7), Inches(2.5), title)
        tb_o = s4.shapes.add_textbox(left + Inches(0.2), top + Inches(0.65), Inches(5.3), Inches(1.7))
        tf_o = tb_o.text_frame
        tf_o.word_wrap = True
        p_o = tf_o.paragraphs[0]
        p_o.text = desc
        p_o.font.size = Pt(12)
        p_o.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 5: SYSTEM ARCHITECTURE
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_slide_background(s5)
    add_header(s5, "Proposed Solution: 3-Tier FlowGrid Architecture", "SYSTEM PIPELINE")

    add_card(s5, Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2), "End-to-End System Pipeline & Technology Stack")
    tb_arch = s5.shapes.add_textbox(Inches(1.1), Inches(2.2), Inches(11.1), Inches(4.4))
    tf_arch = tb_arch.text_frame
    tf_arch.word_wrap = True

    arch_layers = [
        ("Tier 1: Web Dashboard UI (Client Layer)", "Semantic HTML5, responsive dark-glassmorphism CSS3, and vanilla JavaScript. Features dynamic 8-segment road cards, active range tags ([L], [R], In Range), interactive card selection, and real-time ancestor update path highlighting."),
        ("Tier 2: Lightweight HTTP Bridge (server.py)", "Zero external Python dependencies (uses Python standard library http.server, subprocess, json). Serves static dashboard assets, manages persistent corridor state, and validates HTTP REST requests."),
        ("Tier 3: Core Fenwick Tree Engine (C++17 Binary)", "Compiled standalone binary (traffic_analytics.exe). Implements O(n) build, O(log n) updates and prefix/range queries, and congestion classification. Invoked via subprocess CLI JSON mode (--query, --update, --state)."),
        ("Communication Flow", "Browser fetch() -> HTTP REST POST (/api/query, /api/update) -> server.py subprocess -> traffic_analytics.exe execution -> JSON stdout response back to Dashboard."),
        ("Academic Simulation Note", "Operates on simulated sensor volume data across 8 arterial segments for rigorous algorithmic demonstration and validation.")
    ]
    for name, detail in arch_layers:
        p = tf_arch.add_paragraph()
        p.text = "• " + name + ": " + detail
        p.font.size = Pt(11)
        p.font.color.rgb = ACCENT_CYAN if "Tier" in name else (ACCENT_GREEN if "Communication" in name else TEXT_MUTED)
        p.space_after = Pt(6)

    # =========================================================================
    # SLIDE 6: WHY FENWICK TREE (COMPARATIVE TABLE)
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_slide_background(s6)
    add_header(s6, "Why Fenwick Tree? Comparative Data Structure Analysis", "ALGORITHMIC SELECTION")

    table_shape = s6.shapes.add_table(5, 5, Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.0))
    table = table_shape.table

    headers = ["Data Structure", "Point Update", "Prefix Sum", "Range Sum [L..R]", "Auxiliary Space"]
    comp_data = [
        ["Naive Array", "O(1)", "O(n)", "O(n)", "O(1) (No acceleration)"],
        ["Prefix Sum Array", "O(n)", "O(1)", "O(1)", "O(n) (Cascade penalty)"],
        ["Segment Tree", "O(log n)", "O(log n)", "O(log n)", "O(4n) (High pointer overhead)"],
        ["Fenwick Tree (BIT)", "O(log n)", "O(log n)", "O(log n)", "O(n) (Flat cache-friendly array)"]
    ]

    for c, h in enumerate(headers):
        cell = table.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(13)
            p.font.bold = True
            p.font.color.rgb = ACCENT_CYAN

    for r, row in enumerate(comp_data):
        for c, val in enumerate(row):
            cell = table.cell(r + 1, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG if r < 3 else RGBColor(30, 58, 138)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(12)
                p.font.bold = (r == 3)
                p.font.color.rgb = ACCENT_GREEN if (r == 3 and c > 0) else TEXT_LIGHT

    add_card(s6, Inches(0.8), Inches(5.9), Inches(11.7), Inches(1.1), "")
    tb_w = s6.shapes.add_textbox(Inches(1.0), Inches(6.0), Inches(11.3), Inches(0.9))
    tf_w = tb_w.text_frame
    p_w = tf_w.paragraphs[0]
    p_w.text = "Key Takeaway: The Fenwick Tree achieves the optimal Pareto frontier — matching Segment Tree query/update speed (O(log n)) while using 75% less memory (O(n) vs O(4n)), eliminating pointer chasing, and executing on flat CPU caches."
    p_w.font.size = Pt(11)
    p_w.font.color.rgb = TEXT_LIGHT

    # =========================================================================
    # SLIDE 7: FENWICK TREE REPRESENTATION & LOWBIT
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_slide_background(s7)
    add_header(s7, "Fenwick Tree Representation & The lowbit() Concept", "MATHEMATICAL FOUNDATION")

    add_card(s7, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "The lowbit(x) Primitive: Formula & Mechanics")
    tb_l = s7.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.4))
    tf_l = tb_l.text_frame
    tf_l.word_wrap = True
    pts_l = [
        "Mathematical Formula: lowbit(x) = x & (-x)",
        "Two's Complement Property: In two's complement arithmetic, -x = (~x) + 1. Bitwise ANDing x with -x isolates the least significant 1-bit.",
        "Interval of Responsibility: Node bit[i] stores cumulative sum over the contiguous range [i - lowbit(i) + 1 .. i].",
        "Length of Covered Range: Exactly lowbit(i) segments.",
        "Even vs Odd Indices: Odd indices have lowbit = 1 (leaves storing single elements). Powers of 2 cover all preceding segments [1..i]."
    ]
    for pt in pts_l:
        p = tf_l.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    add_card(s7, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "1-Based Responsibility Table (n = 8)")
    tb_m = s7.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.4))
    tf_m = tb_m.text_frame
    tf_m.word_wrap = True
    mappings = [
        "bit[1] : 0001_2 | lowbit=1 | Interval [1..1] | Covers Seg 1 (Leaf)",
        "bit[2] : 0010_2 | lowbit=2 | Interval [1..2] | Covers Segs 1..2 (Sub-root)",
        "bit[3] : 0011_2 | lowbit=1 | Interval [3..3] | Covers Seg 3 (Leaf)",
        "bit[4] : 0100_2 | lowbit=4 | Interval [1..4] | Covers Segs 1..4 (Branch Root)",
        "bit[5] : 0101_2 | lowbit=1 | Interval [5..5] | Covers Seg 5 (Leaf)",
        "bit[6] : 0110_2 | lowbit=2 | Interval [5..6] | Covers Segs 5..6 (Sub-root)",
        "bit[7] : 0111_2 | lowbit=1 | Interval [7..7] | Covers Seg 7 (Leaf)",
        "bit[8] : 1000_2 | lowbit=8 | Interval [1..8] | Covers Segs 1..8 (Global Root)"
    ]
    for mp in mappings:
        p = tf_m.add_paragraph()
        p.text = mp
        p.font.size = Pt(11)
        p.font.color.rgb = ACCENT_CYAN
        p.space_after = Pt(4)

    # =========================================================================
    # SLIDE 8: DYNAMIC POINT UPDATE OPERATION
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_slide_background(s8)
    add_header(s8, "Dynamic Point Update Operation: O(log n)", "CORE ALGORITHMS")

    add_card(s8, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Algorithm: update(idx, delta) & setCount()")
    tb_u = s8.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.4))
    tf_u = tb_u.text_frame
    tf_u.word_wrap = True
    pts_u = [
        "Goal: Add delta to segment idx and propagate the change upward to all ancestor nodes in the implicit Fenwick Tree.",
        "Loop Mechanism: Start at i = idx; while i <= n, add delta to bit[i] and advance to parent: i += lowbit(i).",
        "Tree Ascent: Adding lowbit(i) flips the lowest set bit and carries upward, strictly visiting ancestor nodes in O(log n) steps.",
        "setCount(idx, newCount): Computes delta = newCount - trafficCounts[idx], updates the raw array, then invokes update(idx, delta).",
        "Boundary Validation: Rejects invalid segment indices (idx < 1 or idx > n) and negative vehicle counts."
    ]
    for pt in pts_u:
        p = tf_u.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    add_card(s8, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Trace: Updating Segment 5 by delta = +14")
    tb_tr = s8.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.4))
    tf_tr = tb_tr.text_frame
    tf_tr.word_wrap = True
    traces = [
        "Initial State: trafficCounts[5] = 20 vehicles.",
        "Dynamic Event: Sensor reports 34 vehicles (surge/accident).",
        "Computed Delta: 34 - 20 = +14 vehicles.",
        "",
        "Step 1: i = 5 (0101_2): lowbit(5) = 1.",
        "  bit[5] += 14  -->  20 + 14 = 34. Next: 5 + 1 = 6.",
        "Step 2: i = 6 (0110_2): lowbit(6) = 2.",
        "  bit[6] += 14  -->  25 + 14 = 39. Next: 6 + 2 = 8.",
        "Step 3: i = 8 (1000_2): lowbit(8) = 8.",
        "  bit[8] += 14  -->  87 + 14 = 101. Next: 8 + 8 = 16 > 8.",
        "",
        "Update Path: [5 -> 6 -> 8] (Exactly 3 nodes touched out of 8!)"
    ]
    for tr in traces:
        p = tf_tr.add_paragraph()
        p.text = tr
        p.font.size = Pt(11)
        p.font.color.rgb = ACCENT_GREEN if "Update Path" in tr else (ACCENT_CYAN if "Step" in tr else TEXT_LIGHT)
        p.space_after = Pt(2)

    # =========================================================================
    # SLIDE 9: PREFIX SUM & RANGE QUERY OPERATION
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_slide_background(s9)
    add_header(s9, "Prefix Sum & Corridor Range Aggregation: O(log n)", "CORE ALGORITHMS")

    add_card(s9, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Prefix Sum: prefixSum(idx)")
    tb_ps = s9.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.4))
    tf_ps = tb_ps.text_frame
    tf_ps.word_wrap = True
    pts_ps = [
        "Goal: Compute total traffic from segment 1 to idx: sum(traffic[1..idx]).",
        "Loop Mechanism: Start at i = idx; while i > 0, accumulate bit[i] into sum, and step to left disjoint interval: i -= lowbit(i).",
        "Binary Decomposition: Each step strips the lowest set bit from index i, decomposing the prefix into at most log2(n) canonical sub-intervals.",
        "Time Complexity: At most log2(n) iterations -> strictly O(log n)."
    ]
    for pt in pts_ps:
        p = tf_ps.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    add_card(s9, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Corridor Range Query & Congestion Logic")
    tb_rs = s9.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.4))
    tf_rs = tb_rs.text_frame
    tf_rs.word_wrap = True
    pts_rs = [
        "Principle of Inclusion-Exclusion:",
        "  rangeSum(L, R) = prefixSum(R) - prefixSum(L - 1)",
        "Example Corridor [3..6]:",
        "  prefixSum(6) = bit[6] + bit[4] = 25 + 43 = 68",
        "  prefixSum(2) = bit[2] = 19",
        "  rangeSum(3, 6) = 68 - 19 = 49 vehicles.",
        "",
        "Congestion Classification Rule (Threshold = 55 veh):",
        "  • Volume < 55  -->  NORMAL",
        "  • Volume >= 55 -->  POTENTIALLY CONGESTED",
        "Boundary Equality (Volume = 55, Threshold = 55):",
        "  Classified as POTENTIALLY CONGESTED (>= rule enforced)."
    ]
    for pt in pts_rs:
        p = tf_rs.add_paragraph()
        p.text = pt
        p.font.size = Pt(11)
        p.font.color.rgb = ACCENT_CYAN if "rangeSum(" in pt else (ACCENT_CORAL if "POTENTIALLY" in pt else TEXT_LIGHT)
        p.space_after = Pt(3)

    # =========================================================================
    # SLIDE 10: WORKED EXAMPLE (CASE STUDY WALKTHROUGH)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_slide_background(s10)
    add_header(s10, "Step-by-Step Case Study Walkthrough: Baseline vs Surge", "NUMERICAL VERIFICATION")

    add_card(s10, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Phase 1: Baseline State & Initial Query")
    tb_w1 = s10.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.4))
    tf_w1 = tb_w1.text_frame
    tf_w1.word_wrap = True
    w1_lines = [
        "Initial Traffic Counts [1..8]:",
        "  [12, 7, 15, 9, 20, 5, 11, 8]",
        "Constructed Initial BIT Array [1..8]:",
        "  [12, 19, 15, 43, 20, 25, 11, 87]",
        "Total Corridor Traffic: prefixSum(8) = 87 veh [PASS]",
        "",
        "Query Corridor [3..6] (Threshold = 55):",
        "  prefixSum(6) = bit[6] + bit[4] = 25 + 43 = 68",
        "  prefixSum(2) = bit[2] = 19",
        "  rangeSum(3, 6) = 68 - 19 = 49 vehicles.",
        "",
        "Verdict: 49 < 55 --> NORMAL [VERIFIED]"
    ]
    for ln in w1_lines:
        p = tf_w1.add_paragraph()
        p.text = ln
        p.font.size = Pt(11)
        p.font.color.rgb = ACCENT_GREEN if "[PASS]" in ln or "[VERIFIED]" in ln else TEXT_LIGHT
        p.space_after = Pt(3)

    add_card(s10, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Phase 2: Dynamic Surge & Re-Query")
    tb_w2 = s10.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.4))
    tf_w2 = tb_w2.text_frame
    tf_w2.word_wrap = True
    w2_lines = [
        "Point Update: Segment 5 updated from 20 to 34",
        "Delta: 34 - 20 = +14 vehicles.",
        "Ancestor Traversal Path: bit[5], bit[6], bit[8]",
        "Updated Traffic Counts [1..8]:",
        "  [12, 7, 15, 9, 34, 5, 11, 8]",
        "Updated BIT Array [1..8]:",
        "  [12, 19, 15, 43, 34, 39, 11, 101] [PASS]",
        "",
        "Re-Query Corridor [3..6] (Threshold = 55):",
        "  prefixSum(6) = bit[6] + bit[4] = 39 + 43 = 82",
        "  prefixSum(2) = bit[2] = 19",
        "  rangeSum(3, 6) = 82 - 19 = 63 vehicles.",
        "",
        "Verdict: 63 >= 55 --> POTENTIALLY CONGESTED [VERIFIED]"
    ]
    for ln in w2_lines:
        p = tf_w2.add_paragraph()
        p.text = ln
        p.font.size = Pt(11)
        p.font.color.rgb = ACCENT_CORAL if "CONGESTED" in ln else (ACCENT_GREEN if "[PASS]" in ln else TEXT_LIGHT)
        p.space_after = Pt(3)

    # =========================================================================
    # SLIDE 11: WEB DASHBOARD (MEMBER 4 INTEGRATION)
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_slide_background(s11)
    add_header(s11, "FlowGrid Web Dashboard: Interface & Features", "FRONTEND DASHBOARD")

    add_card(s11, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Visual Interactive Components")
    tb_d1 = s11.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.4))
    tf_d1 = tb_d1.text_frame
    tf_d1.word_wrap = True
    pts_d1 = [
        "8 Road Segment Cards: Real-world naming (North Arterial, Central Expressway, etc.), vehicle count, and capacity progress bar.",
        "Interactive Click-to-Select: Clicking any road segment card automatically selects it in the Point Update panel with a turquoise glow.",
        "Active Range Markers: Segment cards within active corridor [L..R] display distinct [L], [R], and In Range badges.",
        "Live Backend Status Pill: Top navbar displays real-time connection status ('C++ Fenwick Core: Connected')."
    ]
    for pt in pts_d1:
        p = tf_d1.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    add_card(s11, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Structural Inspector & Execution Logs")
    tb_d2 = s11.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.4))
    tf_d2 = tb_d2.text_frame
    tf_d2.word_wrap = True
    pts_d2 = [
        "BIT Array Structural Inspector: Interactive table showing binary notation, lowbit(i), covered interval, mathematical formula, and stored node value.",
        "Live Ancestor Tracing: Ancestor nodes (bit[5], bit[6], bit[8]) highlight across the table and hierarchy view during updates.",
        "Tree Hierarchy Cards: Visual breakdown of Global Root (2^3), Branch Roots (2^2), Sub-roots (2^1), and Leaves (2^0).",
        "Audit Log Feed: Timestamped chronological record of all initialization, update, range query, and classification events."
    ]
    for pt in pts_d2:
        p = tf_d2.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    # =========================================================================
    # SLIDE 12: FULL-STACK INTEGRATION WORKFLOW
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_slide_background(s12)
    add_header(s12, "Full-Stack Integration: Python Server & C++ CLI", "FULL-STACK ARCHITECTURE")

    add_card(s12, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Python HTTP Server (server.py)")
    tb_s1 = s12.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.4))
    tf_s1 = tb_s1.text_frame
    tf_s1.word_wrap = True
    pts_s1 = [
        "Standard Library Architecture: Zero external pip dependencies; uses http.server and subprocess for lightweight local deployment.",
        "Static Asset Delivery: Serves app/index.html, style.css, and script.js seamlessly on port 8000.",
        "State Management: Tracks current 8-segment traffic vector [12, 7, 15, 9, 20, 5, 11, 8] across HTTP requests.",
        "Robust Input Validation: Validates corridor range (1 <= L <= R <= 8) and non-negative counts before spawning subprocess."
    ]
    for pt in pts_s1:
        p = tf_s1.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    add_card(s12, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "REST API Endpoints & C++ CLI Bridge")
    tb_s2 = s12.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.4))
    tf_s2 = tb_s2.text_frame
    tf_s2.word_wrap = True
    api_pts = [
        "GET /api/status: Returns backend health, binary path, and online indicator.",
        "GET /api/state: Queries C++ --state and returns current traffic counts and BIT array.",
        "POST /api/query: Accepts {left, right, threshold}, executes traffic_analytics.exe --query, and returns sum, status, and prefix breakdown.",
        "POST /api/update: Accepts {segment, newCount}, executes traffic_analytics.exe --update, and returns updated arrays and ancestor path.",
        "POST /api/reset: Restores benchmark initial state (87 veh).",
        "C++ CLI Output: Structured JSON emitted directly to stdout."
    ]
    for pt in api_pts:
        p = tf_s2.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = ACCENT_CYAN if "GET" in pt or "POST" in pt else TEXT_LIGHT
        p.space_after = Pt(5)

    # =========================================================================
    # SLIDE 13: COMPLEXITY ANALYSIS & PROOFS
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_slide_background(s13)
    add_header(s13, "Algorithmic Complexity Analysis & Theoretical Proofs", "COMPLEXITY ANALYSIS")

    table_c_shape = s13.shapes.add_table(5, 4, Inches(0.8), Inches(1.6), Inches(11.7), Inches(3.6))
    table_c = table_c_shape.table

    h_c = ["Operation", "Time Complexity", "Auxiliary Space", "Theoretical Proof Summary"]
    d_c = [
        ["Tree Build (buildFenwick)", "O(n)", "O(n)", "Forward-propagation pushes bit[i] to parent i + lowbit(i); exactly n iterations."],
        ["Point Update (update / setCount)", "O(log n)", "O(1)", "Advances via i += lowbit(i); loops at most floor(log2 n) + 1 times."],
        ["Prefix Sum Query (prefixSum)", "O(log n)", "O(1)", "Strips lowest set bit via i -= lowbit(i); loops at most log2 n times."],
        ["Range Query (rangeSum)", "O(log n)", "O(1)", "Difference of two prefix sums: 2 * O(log n) = O(log n)."]
    ]

    for c, h in enumerate(h_c):
        cell = table_c.cell(0, c)
        cell.text = h
        cell.fill.solid()
        cell.fill.fore_color.rgb = CARD_BG
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(12)
            p.font.bold = True
            p.font.color.rgb = ACCENT_CYAN

    for r, row in enumerate(d_c):
        for c, val in enumerate(row):
            cell = table_c.cell(r + 1, c)
            cell.text = val
            cell.fill.solid()
            cell.fill.fore_color.rgb = CARD_BG
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(11)
                p.font.color.rgb = ACCENT_GREEN if c == 1 else TEXT_LIGHT

    add_card(s13, Inches(0.8), Inches(5.5), Inches(11.7), Inches(1.5), "Hardware & CPU Cache Locality Advantages")
    tb_adv = s13.shapes.add_textbox(Inches(1.0), Inches(5.9), Inches(11.3), Inches(1.0))
    tf_adv = tb_adv.text_frame
    pts_adv = [
        "Contiguous Memory Locality: Stored in a single flat array; zero CPU pointer chasing or cache thrashing.",
        "Bitwise Hardware Acceleration: lowbit(x) compiles to single-cycle machine instructions (AND, NEG, BLSI).",
        "Minimal Auxiliary Footprint: Requires only (n + 1) integer slots compared to 4n for Segment Trees."
    ]
    for pt in pts_adv:
        p = tf_adv.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(2)

    # =========================================================================
    # SLIDE 14: TESTING & VERIFICATION (32/32 TESTS)
    # =========================================================================
    s14 = prs.slides.add_slide(blank_layout)
    set_slide_background(s14)
    add_header(s14, "Empirical Verification: 32/32 Automated Tests Passed", "TESTING & VALIDATION")

    add_card(s14, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Automated Test Suite (tests/test_flowgrid.py)")
    tb_t1 = s14.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.4))
    tf_t1 = tb_t1.text_frame
    tf_t1.word_wrap = True
    pts_t1 = [
        "Execution Outcome: 32 PASSED, 0 FAILED (100% Precision).",
        "Test Group 1 (Construction): O(n) forward propagation verified; computed BIT matches expected [12, 19, 15, 43, 20, 25, 11, 87]; total volume = 87 veh.",
        "Test Group 2 (Prefix Sums): Decompositions verified across all indices [1..1] through [1..8].",
        "Test Group 3 (Range Sums): Baseline [3..6] = 49 veh, single segment [5..5] = 20 veh, full corridor [1..8] = 87 veh.",
        "Test Group 4 (Point Updates): Traversal path [5, 6, 8] verified; raw array and BIT updated correctly."
    ]
    for pt in pts_t1:
        p = tf_t1.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = ACCENT_GREEN if "32 PASSED" in pt else TEXT_LIGHT
        p.space_after = Pt(6)

    add_card(s14, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Safety, Boundary & API Integration Tests")
    tb_t2 = s14.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.4))
    tf_t2 = tb_t2.text_frame
    tf_t2.word_wrap = True
    pts_t2 = [
        "Test Group 5 (Post-Update Query): Re-evaluation of [3..6] returns exactly 63 veh.",
        "Test Group 6 (Congestion Logic): Below threshold (49 < 55 -> NORMAL), at boundary (55 == 55 -> POTENTIALLY CONGESTED), above threshold (63 >= 55 -> POTENTIALLY CONGESTED).",
        "Test Group 7 (Boundary Validations): Inverted range [6..3], out-of-bounds indices (10), and negative counts (-10) rejected gracefully.",
        "Test Group 8 (Repeated Updates): 3 sequential updates maintain cumulative consistency (107 veh).",
        "Test Group 9 (HTTP REST API): Validated /api/status, /api/state, /api/query, /api/update, /api/reset.",
        "Test Group 11 (C++ Benchmark): Default argc == 1 non-menu benchmark runs end-to-end."
    ]
    for pt in pts_t2:
        p = tf_t2.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(4)

    # =========================================================================
    # SLIDE 15: LIMITATIONS & FUTURE ROADMAP
    # =========================================================================
    s15 = prs.slides.add_slide(blank_layout)
    set_slide_background(s15)
    add_header(s15, "Limitations & Future Roadmap", "SCALABILITY & ROADMAP")

    add_card(s15, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Current System Limitations")
    tb_l1 = s15.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.4))
    tf_l1 = tb_l1.text_frame
    tf_l1.word_wrap = True
    pts_l1 = [
        "1D Linear Corridor Topology: Current implementation assumes a linear sequence of road segments. Metropolitan networks contain complex graph intersections and multi-arterial grids.",
        "Fixed Segment Array Capacity: Initialized for fixed size n = 8. Dynamically adding road segments requires vector reallocation.",
        "Single Vehicle Count Metric: Tracks aggregate vehicular count only, without factoring vehicle classes (trucks vs cars) or speed.",
        "Point Update Asymmetry: Optimized for point updates. Range updates (adding delta to an entire range) require dual-tree difference encoding."
    ]
    for pt in pts_l1:
        p = tf_l1.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    add_card(s15, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Future Engineering Roadmap")
    tb_l2 = s15.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.4))
    tf_l2 = tb_l2.text_frame
    tf_l2.word_wrap = True
    pts_l2 = [
        "2D Fenwick Tree Extension: Extend to 2D Binary Indexed Trees to monitor spatial M x N road grids and multi-directional intersections in O(log^2 n) time.",
        "Dual-Tree Difference Encoding: Maintain two BITs (D[i] and D[i] * i) to support simultaneous O(log n) range updates and range queries.",
        "Predictive Machine Learning Forecasting: Couple BIT cumulative volume with ARIMA/LSTM time-series models to anticipate congestion 15 minutes ahead.",
        "IoT Telemetry Integration: Ingest physical roadside induction loop sensor feeds via live WebSockets directly into the dashboard."
    ]
    for pt in pts_l2:
        p = tf_l2.add_paragraph()
        p.text = "• " + pt
        p.font.size = Pt(11)
        p.font.color.rgb = ACCENT_CYAN if "Extension" in pt else TEXT_LIGHT
        p.space_after = Pt(8)

    # =========================================================================
    # SLIDE 16: CONCLUSION & ACADEMIC REFERENCES
    # =========================================================================
    s16 = prs.slides.add_slide(blank_layout)
    set_slide_background(s16)
    add_header(s16, "Conclusion, Viva Summary & Academic References", "CONCLUSION & REFERENCES")

    add_card(s16, Inches(0.8), Inches(1.6), Inches(5.6), Inches(5.2), "Key Accomplishments & Evaluation Summary")
    tb_c_end = s16.shapes.add_textbox(Inches(1.0), Inches(2.2), Inches(5.2), Inches(4.4))
    tf_c_end = tb_c_end.text_frame
    tf_c_end.word_wrap = True
    conclusions = [
        "Complete 4-Member Collaboration: Seamlessly synthesizes Member 1 (data), Member 2 (core BIT), Member 3 (queries & demo), and Member 4 (dashboard & server).",
        "Production-Grade Algorithmic Rigor: 32 automated tests verify 100% mathematical and logical parity across C++ and web tiers.",
        "Sub-Linear Operational Speed: O(log n) dynamic updates and range queries executed on flat, cache-friendly arrays.",
        "Standardized Congestion Logic: Strictly enforces Volume < 55 -> NORMAL, Volume >= 55 -> POTENTIALLY CONGESTED.",
        "Academic Viva Ready: Complete documentation, verified test suite, C++ benchmark, and interactive dashboard fully prepared."
    ]
    for c in conclusions:
        p = tf_c_end.add_paragraph()
        p.text = "✓ " + c
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = ACCENT_GREEN
        p.space_after = Pt(8)

    add_card(s16, Inches(6.8), Inches(1.6), Inches(5.7), Inches(5.2), "Academic & Technical References")
    tb_ref = s16.shapes.add_textbox(Inches(7.0), Inches(2.2), Inches(5.3), Inches(4.4))
    tf_ref = tb_ref.text_frame
    tf_ref.word_wrap = True
    references = [
        "1. Peter M. Fenwick (1994). 'A New Data Structure for Cumulative Frequency Tables'. Software: Practice and Experience, 24(3), pp. 327–336.",
        "2. Thomas H. Cormen, Charles E. Leiserson, Ronald L. Rivest, Clifford Stein (2022). 'Introduction to Algorithms' (CLRS), 4th Edition, MIT Press.",
        "3. Donald E. Knuth (1997). 'The Art of Computer Programming, Volume 3: Sorting and Searching', 2nd Edition, Addison-Wesley.",
        "4. ISO/IEC 14882:2017. 'Programming Languages — C++ (C++17 Standard Specification)', International Organization for Standardization.",
        "5. Python Software Foundation. 'Python 3 Standard Library Documentation: http.server and subprocess modules', docs.python.org."
    ]
    for r in references:
        p = tf_ref.add_paragraph()
        p.text = r
        p.font.size = Pt(11)
        p.font.color.rgb = TEXT_LIGHT
        p.space_after = Pt(8)

    # -------------------------------------------------------------------------
    # Save Presentation
    # -------------------------------------------------------------------------
    updated_path = os.path.join("presentation", "Fenwick_Tree_Case_Study_Updated.pptx")
    prs.save(updated_path)
    print(f"[SUCCESS] Updated presentation saved to: {updated_path}")

    # Also attempt to update the original file if writable
    original_path = os.path.join("presentation", "Fenwick_Tree_Case_Study.pptx")
    try:
        prs.save(original_path)
        print(f"[SUCCESS] Original presentation also updated at: {original_path}")
    except PermissionError:
        print(f"[NOTE] Original file {original_path} is currently locked by PowerPoint. The updated version is safely preserved at {updated_path}.")

    return updated_path


if __name__ == "__main__":
    create_deck()
